"""
helpers.py — общие вспомогательные функции для тестов (синтетические данные).
"""
from __future__ import annotations

from datetime import date

import pandas as pd

import history_reader as hr


# --- История RSV и таймлайн --------------------------------------------------------
def make_history(rows):
    cols = hr.RSV_DATA_COLUMNS + ["file_year", "file_month"]
    df = pd.DataFrame(rows).reindex(columns=cols)
    for c in ("tn", "vop", "vred_dni", "otp_dni", "otp_bud_m", "po_sredn", "tar1", "kol_rd"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def vr(month, vred, cls, tar, tn=1, year=2025, kol=20):
    return dict(tn=tn, vop=49, vred_dni=vred, kol_rd=kol, bal_vredn=cls, tar1=tar, file_year=year, file_month=month)


def year_rows(skip=(), vred=20, cls=3.1, tar=0.29):
    return [vr(m, vred, cls, tar) for m in range(1, 13) if m not in skip]


def leave_of(rows, **kw):
    tl = hr.build_employee_timeline(1, make_history(rows), hire_date=date(2025, 1, 10))
    blocks = hr.build_work_year_blocks(tl, **kw)
    return tl, blocks[0]["leave"]


def by_class(leave):
    return {c["klass"]: c for c in leave["classes"]}


# --- Файлы RSV на диске ---------------------------------------------------------------
def write_rsv(path, rows, header_rows=2):
    """rows — список словарей {колонка RSV: значение}; данные с 3-й строки."""
    from openpyxl import Workbook

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    for r in range(1, header_rows + 1):
        ws.cell(r, 1, "шапка")
    for i, rec_ in enumerate(rows):
        for j, col in enumerate(hr.RSV_DATA_COLUMNS):
            if col in rec_ and rec_[col] is not None:
                ws.cell(3 + i, j + 1, rec_[col])
    wb.save(path)


def rec(tn, **kw):
    base = dict(tn=tn, vop=49, vred_dni=20, kol_rd=20, bal_vredn=3.1)
    base.update(kw)
    return base
