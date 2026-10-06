"""
test_stage6.py — проверки качества данных (синтетика).
Запуск: pytest test_stage6.py -v
"""
from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import Workbook

import history_reader as hr
import insalubrity as ins


def write_rsv(path, rows, header_rows=2):
    """rows — список словарей {колонка RSV: значение}; данные с 3-й строки."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    for r in range(1, header_rows + 1):
        ws.cell(r, 1, "шапка")
    for i, rec in enumerate(rows):
        for j, col in enumerate(hr.RSV_DATA_COLUMNS):
            if col in rec and rec[col] is not None:
                ws.cell(3 + i, j + 1, rec[col])
    wb.save(path)


def rec(tn, **kw):
    base = dict(tn=tn, vop=49, vred_dni=20, kol_rd=20, bal_vredn=3.1)
    base.update(kw)
    return base


BLANK = {}


# --- пустая строка в середине данных ---------------------------------------
def test_blank_row_in_middle_no_longer_truncates(tmp_path):
    f = tmp_path / "2025" / "rsv_01.xlsx"
    write_rsv(f, [rec(1), rec(2), BLANK, rec(3), rec(4)])
    logs = []
    df = hr.read_rsv_file(f, log=logs.append)
    assert list(df["tn"]) == [1, 2, 3, 4]
    assert any("пустая строка 5" in m and "2 строк" in m for m in logs)


def test_clean_file_has_no_warnings(tmp_path):
    f = tmp_path / "2025" / "rsv_01.xlsx"
    write_rsv(f, [rec(1), rec(2)])
    logs = []
    assert len(hr.read_rsv_file(f, log=logs.append)) == 2
    assert logs == []


def test_trailing_blank_rows_are_silent(tmp_path):
    f = tmp_path / "2025" / "rsv_01.xlsx"
    write_rsv(f, [rec(1), rec(2), BLANK, BLANK])
    logs = []
    assert len(hr.read_rsv_file(f, log=logs.append)) == 2
    assert logs == []


def test_totals_row_stops_reading_and_warns_about_rows_below(tmp_path):
    f = tmp_path / "2025" / "rsv_01.xlsx"
    totals = dict(vred_dni=40)                  # итоговая строка: нет tn, но есть сумма
    write_rsv(f, [rec(1), rec(2), totals, rec(3)])
    logs = []
    df = hr.read_rsv_file(f, log=logs.append)
    assert list(df["tn"]) == [1, 2]
    assert any("после итоговой строки" in m and "НЕ прочитаны" in m for m in logs)


def test_totals_row_alone_is_not_an_error(tmp_path):
    f = tmp_path / "2025" / "rsv_01.xlsx"
    write_rsv(f, [rec(1), rec(2), dict(vred_dni=40)])
    logs = []
    assert len(hr.read_rsv_file(f, log=logs.append)) == 2
    assert logs == []


def test_read_rsv_file_still_works_without_log(tmp_path):
    f = tmp_path / "2025" / "rsv_01.xlsx"
    write_rsv(f, [rec(1), BLANK, rec(2)])
    assert list(hr.read_rsv_file(f)["tn"]) == [1, 2]


# --- табельные номера и числа в истории ---------------------------------------
def test_tn_types_are_normalized_in_history(tmp_path):
    write_rsv(tmp_path / "2025" / "rsv_01.xlsx", [rec(5), rec("5"), rec(5.0), rec(" 5 ")])
    logs = []
    h = hr.build_history(tmp_path, log=logs.append)
    assert list(h["tn"]) == [5, 5, 5, 5]
    assert not any("tn" in m for m in logs)


def test_unreadable_tn_and_numbers_are_reported_not_silent(tmp_path):
    write_rsv(tmp_path / "2025" / "rsv_01.xlsx", [rec(5), rec("А12"), rec(6, vred_dni="много")])
    logs = []
    h = hr.build_history(tmp_path, log=logs.append)
    assert h["tn"].isna().sum() == 1
    assert any("«tn»" in m and "А12" in m and "01.2025" in m for m in logs)
    assert any("«vred_dni»" in m and "много" in m for m in logs)


# --- конфликтующие дубли в справочниках ------------------------------------------
def test_conflicting_duplicates_are_logged():
    lschet = pd.DataFrame({"tn": [1, 1, 2], "fio": ["А", "Б", "В"], "pr_dn": [1, 1, 2]})
    logs = []
    n = ins.warn_conflicting_duplicates(lschet, ["tn"], ["fio", "pr_dn"], "lschet.dbf", keep="last", log=logs.append)
    assert n == 1 and "lschet.dbf" in logs[0] and "1" in logs[0] and "последняя" in logs[0]


def test_identical_duplicates_are_silent():
    kal = pd.DataFrame({"pr_dn": [1, 1], "god": [2026, 2026], "mes": [5, 5], "kol_rd": [21, 21], "fond_ch": [168, 168]})
    logs = []
    n = ins.warn_conflicting_duplicates(kal, ["pr_dn", "god", "mes"], ["kol_rd", "fond_ch"],
                                        "kalend.dbf", keep="first", log=logs.append)
    assert n == 0 and logs == []


def test_conflicting_calendar_rows_reported_with_first_kept():
    kal = pd.DataFrame({"pr_dn": [1, 1], "god": [2026, 2026], "mes": [5, 5], "kol_rd": [21, 22], "fond_ch": [168, 176]})
    logs = []
    assert ins.warn_conflicting_duplicates(kal, ["pr_dn", "god", "mes"], ["kol_rd", "fond_ch"],
                                           "kalend.dbf", keep="first", log=logs.append) == 1
    assert "первая" in logs[0]


def test_no_duplicates_no_output():
    lschet = pd.DataFrame({"tn": [1, 2], "fio": ["А", "Б"], "pr_dn": [1, 1]})
    logs = []
    assert ins.warn_conflicting_duplicates(lschet, ["tn"], ["fio", "pr_dn"], "lschet.dbf", keep="last", log=logs.append) == 0
    assert logs == []
