"""
test_patches.py
===============
Юнит-тесты на исправленные баги (синтетические данные, реальные DBF/RSV не нужны).

Запуск:
    python test_patches.py
    pytest test_patches.py -v
"""

from __future__ import annotations

import json
import tempfile
import traceback
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import app_env
import history_reader as hr
import insalubrity


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def make_history(rows: list[dict]) -> pd.DataFrame:
    cols = hr.RSV_DATA_COLUMNS + ["file_year", "file_month"]
    df = pd.DataFrame(rows).reindex(columns=cols)
    for c in ("tn", "vop", "vred_dni", "otp_dni", "otp_bud_m", "po_sredn", "tar1", "kol_rd"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def tmpdir() -> Path:
    return Path(tempfile.mkdtemp(prefix="insal_test_"))


# ---------------------------------------------------------------------------
# п.1 — путь RSV и совместимость с «Историей»
# ---------------------------------------------------------------------------
def test_rsv_output_path_has_year_folder():
    p = insalubrity.rsv_output_path(Path("out"), 2026, 5)
    assert p == Path("out") / "2026" / "rsv_05.xlsx"


def test_two_years_same_month_do_not_collide_and_history_sees_them():
    out = tmpdir()
    for year in (2025, 2026):
        path = insalubrity.rsv_output_path(out, year, 5)
        path.parent.mkdir(parents=True, exist_ok=True)
        df = pd.DataFrame([{"tn": 1, "fio": "X", "vop": 49, "vred_dni": year - 2000}])
        insalubrity.save_rsv_excel(df, path, 5, year, 0.29, 0.41, 0.58)

    found = hr.find_rsv_files(out, log=lambda *_: None)
    assert [(y, m) for y, m, _ in found] == [(2025, 5), (2026, 5)]
    values = [hr.read_rsv_file(p).iloc[0]["vred_dni"] for _, _, p in found]
    assert values == [25, 26]


def test_empty_rsv_has_no_broken_sum_formula():
    path = tmpdir() / "rsv_01.xlsx"
    insalubrity.save_rsv_excel(pd.DataFrame(), path, 1, 2026, 0.29, 0.41, 0.58)
    ws = load_workbook(path).active
    assert all(c.value is None for c in ws[3]), "формула SUM(N3:N2) не должна появляться"


# ---------------------------------------------------------------------------
# п.2 — .env: слияние ключей
# ---------------------------------------------------------------------------
def test_env_save_keeps_foreign_keys():
    old = app_env.ENV_FILE_PATH
    app_env.ENV_FILE_PATH = tmpdir() / ".env"
    try:
        app_env.ENV_FILE_PATH.write_text(
            "OUT_DIR=x\nLEAVE_DAYS_31=5\nLEAVE_DAYS_32=8\n", encoding="utf-8"
        )
        app_env.save_env_vars({"CALC_MONTH": "6", "OUT_DIR": "y"})  # как это делает калькулятор
        cfg = app_env.load_env_vars()
        assert cfg["LEAVE_DAYS_31"] == "5" and cfg["LEAVE_DAYS_32"] == "8"
        assert cfg["CALC_MONTH"] == "6" and cfg["OUT_DIR"] == "y"
    finally:
        app_env.ENV_FILE_PATH = old


# ---------------------------------------------------------------------------
# п.3 — JSON правок: битый файл не перезаписывается
# ---------------------------------------------------------------------------
def test_corrupt_override_file_is_not_overwritten():
    p = tmpdir() / "hire.json"
    p.write_text('{"8024": "1988-01-06"', encoding="utf-8")  # оборван
    before = p.read_text(encoding="utf-8")
    try:
        hr.save_hire_date_override(1, "2020-01-01", path=p)
    except hr.OverrideFileError:
        pass
    else:
        raise AssertionError("ожидался OverrideFileError")
    assert p.read_text(encoding="utf-8") == before, "битый файл не должен быть изменён"


def test_override_save_makes_backup_and_is_valid_json():
    p = tmpdir() / "hire.json"
    hr.save_hire_date_override(1, "2020-01-01", path=p)
    hr.save_hire_date_override(2, "2021-02-02", path=p)
    assert json.loads(p.read_text(encoding="utf-8")) == {"1": "2020-01-01", "2": "2021-02-02"}
    bak = p.with_name(p.name + ".bak")
    assert json.loads(bak.read_text(encoding="utf-8")) == {"1": "2020-01-01"}
    assert not list(p.parent.glob("*.tmp")), "временные файлы должны убираться"


def test_empty_override_file_is_treated_as_no_overrides():
    p = tmpdir() / "m.json"
    p.write_text("", encoding="utf-8")
    assert hr.load_month_overrides(p) == {}
    hr.save_month_override(5, 2026, 9, "vred_dni", 22, path=p)
    assert hr.load_month_overrides(p) == {"5": {"2026-09": {"vred_dni": 22}}}


# ---------------------------------------------------------------------------
# п.8 — даты приёма
# ---------------------------------------------------------------------------
def _employees(dnepr):
    return pd.DataFrame([{"tn": 1, "fio": "X", "dnepr": dnepr}])


def test_implausible_dnepr_is_rejected():
    assert hr.get_hire_date(1, _employees(date(1899, 12, 30)), {}) == (None, "dnepr_invalid")
    assert hr.get_hire_date(1, _employees(date(2999, 1, 1)), {}) == (None, "dnepr_invalid")
    assert hr.get_hire_date(1, _employees(date(2015, 3, 2)), {}) == (date(2015, 3, 2), "dnepr")


def test_override_wins_and_bad_override_falls_back():
    emp = _employees(date(2015, 3, 2))
    assert hr.get_hire_date(1, emp, {"1": "2019-05-05"}) == (date(2019, 5, 5), "override")
    assert hr.get_hire_date(1, emp, {"1": "1800-01-01"}) == (date(2015, 3, 2), "dnepr")


def test_save_hire_date_rejects_garbage():
    p = tmpdir() / "h.json"
    for bad in ("1899-12-30", "2999-01-01", "not a date"):
        try:
            hr.save_hire_date_override(1, bad, path=p)
        except ValueError:
            continue
        raise AssertionError(f"{bad!r} должно быть отклонено")
    assert not p.exists()


def test_data_before_hire_is_reported():
    rows = [dict(tn=4, vop=49, vred_dni=20, kol_rd=20, bal_vredn=3.1, tar1=0.29,
                 file_year=2025, file_month=m) for m in (1, 2, 3, 4)]
    h = make_history(rows)
    assert hr.find_data_before_hire(4, h, date(2025, 3, 15)) == [(2025, 1), (2025, 2)]
    assert hr.find_data_before_hire(4, h, date(2024, 1, 1)) == []
    assert hr.find_data_before_hire(4, h, None) == []


# ---------------------------------------------------------------------------
# п.6 — колонка «Тариф» только из строк вредности
# ---------------------------------------------------------------------------
def test_tar1_ignores_vacation_rows():
    rows = [
        dict(tn=1, vop=49, vred_dni=15, kol_rd=21, bal_vredn=3.2, tar1=0.41, file_year=2025, file_month=7),
        dict(tn=1, vop=80, otp_dni=6, kol_rd=21, tar1=1500.0, file_year=2025, file_month=7),
        dict(tn=1, vop=80, otp_dni=14, kol_rd=21, tar1=1600.0, file_year=2025, file_month=8),  # только отпуск
    ]
    tl = hr.build_employee_timeline(1, make_history(rows), hire_date=date(2025, 7, 1))
    by_month = tl.set_index("month")["tar1"]
    assert abs(by_month[7] - 0.41) < 1e-9
    assert pd.isna(by_month[8]), "в месяце только с отпуском тарифа вредности быть не должно"


def test_tar1_fallback_when_vop_unknown():
    rows = [dict(tn=1, vop=None, vred_dni=15, kol_rd=21, bal_vredn=3.2, tar1=0.41,
                 file_year=2025, file_month=7)]
    tl = hr.build_employee_timeline(1, make_history(rows), hire_date=date(2025, 7, 1))
    assert abs(tl.iloc[0]["tar1"] - 0.41) < 1e-9


# ---------------------------------------------------------------------------
# NaN в kol_rd не превращает итоги в NaN
# ---------------------------------------------------------------------------
def test_nan_kol_rd_does_not_poison_totals():
    rows = [dict(tn=3, vop=49, vred_dni=20, kol_rd=None, bal_vredn=3.1, tar1=0.29,
                 file_year=2026, file_month=m) for m in (3, 4)]
    tl = hr.build_employee_timeline(3, make_history(rows), hire_date=date(2026, 3, 1))
    assert not tl["kol_rd"].isna().any()
    blocks = hr.build_work_year_blocks(tl)
    assert blocks and blocks[0]["kol_rd_total"] == blocks[0]["kol_rd_total"]  # не NaN


# ---------------------------------------------------------------------------
# п.5 — потолок «вышло» и признак смешения единиц
# ---------------------------------------------------------------------------
def _year_rows(tn=1, year=2025, vred=22, kol_rd=21, extra=()):
    rows = [dict(tn=tn, vop=49, vred_dni=vred, kol_rd=kol_rd, bal_vredn=3.2, tar1=0.41,
                 file_year=year, file_month=m) for m in range(1, 13)]
    rows.extend(extra)
    return rows


def test_leave_is_capped_by_norm_and_flagged():
    extra = [dict(tn=1, vop=80, otp_dni=28, kol_rd=21, tar1=1500.0, file_year=2025, file_month=7)]
    tl = hr.build_employee_timeline(1, make_history(_year_rows(extra=extra)), hire_date=date(2025, 1, 10))
    leave = hr.build_work_year_blocks(tl)[0]["leave"]
    cls = leave["classes"][0]
    assert cls["vyshlo_raw"] > cls["norm"], "сценарий должен давать превышение"
    assert cls["vyshlo"] == cls["norm"] == 7 and cls["capped"]
    assert leave["vyshlo_total"] <= leave["polozheno_total"]
    assert leave["days_exceed_norm"] is True or leave["days_exceed_norm"] == True  # noqa: E712


def test_leave_not_capped_when_within_norm():
    tl = hr.build_employee_timeline(1, make_history(_year_rows(vred=10)), hire_date=date(2025, 1, 10))
    leave = hr.build_work_year_blocks(tl)[0]["leave"]
    cls = leave["classes"][0]
    assert not cls["capped"] and abs(cls["vyshlo"] - cls["vyshlo_raw"]) < 1e-12
    assert not leave["days_exceed_norm"]


# ---------------------------------------------------------------------------
# п.7 — «нет файла» против «нет записей у сотрудника»
# ---------------------------------------------------------------------------
def test_missing_file_month_is_distinguished_from_no_records():
    rows = []
    for m in (1, 2, 4, 5):  # файла за март нет вообще
        rows.append(dict(tn=1, vop=49, vred_dni=20, kol_rd=20, bal_vredn=3.1, tar1=0.29,
                         file_year=2026, file_month=m))
    # у сотрудника 2 в апреле запись есть, а у сотрудника 1 в апреле — нет (файл при этом есть)
    rows = [r for r in rows if not (r["file_month"] == 4)]
    rows.append(dict(tn=2, vop=49, vred_dni=20, kol_rd=20, bal_vredn=3.1, tar1=0.29,
                     file_year=2026, file_month=4))
    tl = hr.build_employee_timeline(1, make_history(rows), hire_date=date(2026, 1, 1)).set_index("month")
    assert bool(tl.loc[3, "no_file"]) is True      # файла нет
    assert bool(tl.loc[4, "no_file"]) is False     # файл есть, записей у человека нет
    assert bool(tl.loc[4, "has_data"]) is False
    assert bool(tl.loc[1, "no_file"]) is False
    assert "нет" in tl.loc[3, "projection_note"]


def test_blocks_report_missing_file_months():
    rows = [dict(tn=1, vop=49, vred_dni=20, kol_rd=20, bal_vredn=3.1, tar1=0.29,
                 file_year=2026, file_month=m) for m in (1, 2, 4)]
    tl = hr.build_employee_timeline(1, make_history(rows), hire_date=date(2026, 1, 1))
    block = hr.build_work_year_blocks(tl)[0]
    assert block["missing_file_months"] == [(2026, 3)]


def test_empty_file_period_counts_as_existing():
    rows = [dict(tn=1, vop=49, vred_dni=20, kol_rd=20, bal_vredn=3.1, tar1=0.29,
                 file_year=2026, file_month=1)]
    h = make_history(rows)
    h.attrs["periods"] = [(2026, 1), (2026, 2)]  # февральский файл прочитан, но строк в нём нет
    tl = hr.build_employee_timeline(1, h, hire_date=date(2026, 1, 1)).set_index("month")
    assert 2 in tl.index and bool(tl.loc[2, "no_file"]) is False


def test_build_history_records_read_periods():
    out = tmpdir()
    for m in (1, 2):
        p = insalubrity.rsv_output_path(out, 2026, m)
        p.parent.mkdir(parents=True, exist_ok=True)
        df = pd.DataFrame([{"tn": 1, "fio": "X", "vop": 49, "vred_dni": 5}]) if m == 1 else pd.DataFrame()
        insalubrity.save_rsv_excel(df, p, m, 2026, 0.29, 0.41, 0.58)
    h = hr.build_history(out, log=lambda *_: None)
    assert hr.available_periods(h) == {(2026, 1), (2026, 2)}


# ---------------------------------------------------------------------------
# п.4 — устаревшие ручные правки
# ---------------------------------------------------------------------------
def _timeline_sep(vred_in_rsv):
    rows = [dict(tn=5, vop=49, vred_dni=vred_in_rsv, kol_rd=21, bal_vredn=3.1, tar1=0.29,
                 file_year=2026, file_month=m) for m in range(1, 10)]
    return hr.build_employee_timeline(5, make_history(rows), hire_date=date(2026, 1, 1))


def _sep_row(tn_overrides, vred_in_rsv=20):
    blocks = hr.build_work_year_blocks(_timeline_sep(vred_in_rsv), month_overrides=tn_overrides)
    return [x for x in blocks[0]["month_rows"] if x["month"] == 9][0]


def test_legacy_override_conflicting_with_new_rsv_is_stale():
    r = _sep_row({"2026-09": {"vred_dni": 22}})           # старая запись без маркера
    assert r["vred_dni"] == 22                            # приоритет кадровика сохраняется
    assert set(r["stale_fields"]) == {"vred_dni"}
    assert r["stale_fields"]["vred_dni"] == (20, 22)


def test_override_made_on_real_data_is_not_stale():
    r = _sep_row({"2026-09": {"vred_dni": 22, hr.HAD_DATA_KEY: True}})
    assert r["vred_dni"] == 22 and not r["stale_fields"]


def test_override_equal_to_rsv_is_not_stale():
    r = _sep_row({"2026-09": {"vred_dni": 20}})
    assert not r["stale_fields"]


def test_save_override_stores_marker_and_clear_removes_entry():
    p = tmpdir() / "m.json"
    hr.save_month_override(5, 2026, 9, "vred_dni", 22, path=p, had_data=False)
    assert hr.load_month_overrides(p) == {"5": {"2026-09": {"vred_dni": 22, hr.HAD_DATA_KEY: False}}}
    hr.save_month_override(5, 2026, 9, "vred_dni", 22, path=p, had_data=True)  # подтверждение
    assert hr.load_month_overrides(p)["5"]["2026-09"][hr.HAD_DATA_KEY] is True
    hr.clear_month_override(5, 2026, 9, "vred_dni", path=p)
    assert hr.load_month_overrides(p) == {}, "маркер не должен переживать последнюю правку"


# ---------------------------------------------------------------------------
# пути .env / JSON-правок — относительно папки программы, а не cwd
# ---------------------------------------------------------------------------
def test_paths_are_relative_to_app_dir_not_cwd():
    assert app_env.ENV_FILE_PATH.parent == app_env.data_dir()
    assert hr.DEFAULT_OVERRIDE_FILE.parent == app_env.data_dir()
    assert hr.DEFAULT_MONTH_OVERRIDE_FILE.parent == app_env.data_dir()
    for p in (app_env.ENV_FILE_PATH, hr.DEFAULT_OVERRIDE_FILE, hr.DEFAULT_MONTH_OVERRIDE_FILE):
        assert p.is_absolute(), f"{p} должен быть абсолютным, а не завязанным на текущую папку"


# ---------------------------------------------------------------------------
# insalubrity.py — округление дней отпуска
# ---------------------------------------------------------------------------
def test_fractional_leave_days_are_rounded_not_truncated():
    # 0.5 -> fox_round даёт 1, а не 0, как было бы при astype(int)
    assert insalubrity.fox_round(0.5) == 1
    assert insalubrity.fox_round(2.5) == 3


# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# Увольнение (DATA_UVL) — таймлайн и проекция не продлеваются дальше него
# ---------------------------------------------------------------------------
def _emp_with_uvl(dnepr, data_uvl):
    return pd.DataFrame([{"tn": 1, "fio": "X", "dnepr": dnepr, "data_uvl": data_uvl}])


def test_termination_date_read_and_validated():
    assert hr.get_termination_date(1, _emp_with_uvl(date(2020, 1, 1), date(2026, 8, 31)), date(2020, 1, 1)) \
        == (date(2026, 8, 31), "data_uvl")
    # неправдоподобная (слишком далеко в будущем) -> игнорируется
    d, src = hr.get_termination_date(1, _emp_with_uvl(date(2020, 1, 1), date(2200, 1, 1)), date(2020, 1, 1))
    assert d is None and src == "data_uvl_invalid"
    # раньше даты приёма (переприём) -> игнорируется
    d, src = hr.get_termination_date(1, _emp_with_uvl(date(2020, 1, 1), date(2015, 1, 1)), date(2020, 1, 1))
    assert d is None and src == "data_uvl_before_hire"
    # нет DATA_UVL -> действующий сотрудник
    assert hr.get_termination_date(1, _emp_with_uvl(date(2020, 1, 1), None), date(2020, 1, 1)) == (None, None)


def test_timeline_stops_at_termination_month():
    rows = [dict(tn=1, vop=49, vred_dni=12, kol_rd=20, bal_vredn=3.1, tar1=0.29,
                 file_year=2026, file_month=m) for m in (5, 6, 7, 8)]
    h = make_history(rows)
    h.attrs["periods"] = [(2026, m) for m in (5, 6, 7, 8)]
    tl = hr.build_employee_timeline(1, h, hire_date=date(2026, 5, 26), termination_date=date(2026, 8, 31))
    assert (int(tl["year"].max()), int(tl.loc[tl.year == tl.year.max(), "month"].max())) == (2026, 8)


def test_work_year_block_does_not_project_past_termination():
    rows = [dict(tn=1, vop=49, vred_dni=12, kol_rd=20, bal_vredn=3.1, tar1=0.29,
                 file_year=2026, file_month=m) for m in (5, 6, 7, 8)]
    h = make_history(rows)
    h.attrs["periods"] = [(2026, m) for m in (5, 6, 7, 8)]
    hire_date, term_date = date(2026, 5, 26), date(2026, 8, 31)
    tl = hr.build_employee_timeline(1, h, hire_date=hire_date, termination_date=term_date)
    blocks = hr.build_work_year_blocks(tl, termination_date=term_date)
    assert len(blocks) == 1
    rows_out = blocks[0]["month_rows"]
    assert len(rows_out) == 4, "не должно быть строк после месяца увольнения"
    assert not any(r["is_projected"] for r in rows_out), "после увольнения проекции быть не должно"
    assert max((r["year"], r["month"]) for r in rows_out) == (2026, 8)


def test_real_rsv_data_after_termination_is_not_dropped():
    # финальный расчёт мог попасть в файл следующего месяца — такую РЕАЛЬНУЮ
    # запись отбрасывать нельзя, даже если она формально "после увольнения".
    rows = [
        dict(tn=1, vop=49, vred_dni=12, kol_rd=20, bal_vredn=3.1, tar1=0.29, file_year=2026, file_month=7),
        dict(tn=1, vop=49, vred_dni=5, kol_rd=20, bal_vredn=3.1, tar1=0.29, file_year=2026, file_month=8),
        dict(tn=1, vop=6, po_sredn=2, kol_rd=20, tar1=0.29, file_year=2026, file_month=9),  # расчёт "по среднему" задним числом
    ]
    h = make_history(rows)
    h.attrs["periods"] = [(2026, 7), (2026, 8), (2026, 9)]
    tl = hr.build_employee_timeline(1, h, hire_date=date(2026, 7, 1), termination_date=date(2026, 7, 31))
    assert (int(tl["year"].max()), int(tl.loc[tl.year == tl.year.max(), "month"].max())) == (2026, 9)
    assert bool(tl.loc[(tl.year == 2026) & (tl.month == 9), "has_data"].iloc[0])


def test_active_employee_without_uvl_still_gets_projected_forward():
    # регрессия: у действующего сотрудника (нет DATA_UVL) поведение не должно измениться.
    rows = [dict(tn=1, vop=49, vred_dni=12, kol_rd=20, bal_vredn=3.1, tar1=0.29,
                 file_year=2026, file_month=m) for m in range(1, 9)]
    h = make_history(rows)
    h.attrs["periods"] = [(2026, m) for m in range(1, 9)]
    tl = hr.build_employee_timeline(1, h, hire_date=date(2026, 1, 1), termination_date=None)
    blocks = hr.build_work_year_blocks(tl, termination_date=None)
    assert len(blocks[0]["month_rows"]) == 12
    assert sum(1 for r in blocks[0]["month_rows"] if r["is_projected"]) == 4


# ---------------------------------------------------------------------------
if __name__ == "__main__":
    failed = 0
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    for name, fn in tests:
        try:
            fn()
            print(f"[OK]   {name}")
        except Exception:
            failed += 1
            print(f"[FAIL] {name}\n{traceback.format_exc()}")
    print(f"\n{len(tests) - failed}/{len(tests)} прошло")
    raise SystemExit(1 if failed else 0)
