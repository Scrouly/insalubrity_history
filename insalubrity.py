"""
insalubrity.py
==============
Полный перенос логики VREDN_01.PRG (Clipper/FoxPro) на Python.
Формирует файл Excel (RSV) со специализированной шапкой,
итоговой формулой и автоподбором ширины колонок по содержимому.

Зависимости:
    pip install dbfread pandas openpyxl

Запуск (пример):
    python insalubrity.py --data-dir ./dbf --month 5 --year 2026 --stavka1 0.29 --stavka2 0.41 --stavka3 0.58
"""

from __future__ import annotations

import argparse
import math
import traceback
from pathlib import Path

import pandas as pd
from dbfread import DBF
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Коды операций (vop)
# ---------------------------------------------------------------------------
VOP_VREDNOST = {49}
VOP_OTPUSK = {80, 208, 57}
VOP_OTPUSK_BUD = {81, 214, 116}
VOP_SREDNIY = {6}
VOP_RSV = VOP_VREDNOST | VOP_OTPUSK | VOP_OTPUSK_BUD | VOP_SREDNIY

MONTHS_RU = [
    "", "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"
]

STAVKA_TOL = 1e-4


def fox_round(val: float) -> int:
    """Округление по правилам FoxPro/Clipper (к ближайшему целому)."""
    return int(math.floor(val + 0.5))


def auto_fit_columns(ws, min_width: int = 6, padding: int = 3):
    """Выравнивает ширину колонок листа Excel по максимальной длине содержимого."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                val_str = str(cell.value)
                # Игнорируем длинный текст формулы для расчета ширины (берем стандартную ширину)
                if val_str.startswith("="):
                    val_len = 8
                else:
                    lines = val_str.split("\n")
                    val_len = max(len(line) for line in lines)
                if val_len > max_len:
                    max_len = val_len
        ws.column_dimensions[col_letter].width = max(max_len + padding, min_width)


def load_dbf(path: Path, encoding: str = "cp866") -> pd.DataFrame:
    """Читает .dbf в DataFrame, пропуская удалённые записи."""
    table = DBF(str(path), encoding=encoding, load=True, ignore_missing_memofile=True)
    df = pd.DataFrame(iter(table))
    df.columns = [c.lower() for c in df.columns]
    return df


def find_table(data_dir: Path, base_name: str) -> Path:
    """Ищет таблицу <base_name>.dbf без учёта регистра имени и расширения
    (lschet.dbf, LSCHET.DBF, Lschet.Dbf ...). На Windows файловая система и так
    нечувствительна к регистру, но на Linux/сетевых шарах — нет."""
    wanted = f"{base_name}.dbf".lower()
    try:
        for p in Path(data_dir).iterdir():
            if p.name.lower() == wanted and p.is_file():
                return p
    except OSError:
        pass
    raise FileNotFoundError(f"Не найдена таблица '{base_name}' в {data_dir}")


def classify_bal_vredn(tar1: float, stavka1: float, stavka2: float, stavka3: float):
    """Классификация балла вредности (3.1 / 3.2 / 3.3)."""
    if tar1 is None or (isinstance(tar1, float) and math.isnan(tar1)):
        return None
    for stavka, klass in ((stavka1, 3.1), (stavka2, 3.2), (stavka3, 3.3)):
        if math.isclose(tar1, stavka, abs_tol=STAVKA_TOL):
            return klass
    return None


def rsv_output_path(out_dir: Path, year: int, month: int) -> Path:
    """Путь итогового файла: out_dir/<год>/rsv_ММ.xlsx.

    Именно такую раскладку ищет history_reader.find_rsv_files. Раньше файл
    клался прямо в out_dir, из-за чего (а) «История» его не видела и (б) один и
    тот же месяц разных лет перезаписывал друг друга.
    """
    return out_dir / f"{year:04d}" / f"rsv_{month:02d}.xlsx"


def save_rsv_excel(df_rsv: pd.DataFrame, out_path: Path, month: int, year: int, s1: float, s2: float, s3: float):
    """Сохраняет файл RSV с двухстрочной шапкой, формулой и автоподбором колонок."""
    wb = Workbook()
    ws = wb.active

    # Строка 1: Период и ставки
    period_str = f"{MONTHS_RU[month]}  {year} г."
    ws.cell(row=1, column=10, value=period_str)
    ws.cell(row=1, column=17, value=s1)
    ws.cell(row=1, column=18, value=s2)
    ws.cell(row=1, column=19, value=s3)

    # Строка 2: Названия колонок
    ws.cell(row=2, column=14, value="Вредн")
    ws.cell(row=2, column=15, value="Отпуск")
    ws.cell(row=2, column=16, value="О/б/м")
    ws.cell(row=2, column=17, value="По средн")
    ws.cell(row=2, column=19, value="Балл вредности")
    ws.cell(row=2, column=20, value="Год")
    ws.cell(row=2, column=21, value="Месяц")

    # Порядок колонок данных (21 колонка)
    data_cols = [
        "no", "tn", "fio", "vop", "vf", "sux", "dolgn", "empty_col", "name",
        "pr_dn", "fond_ch", "kol_rd", "sr_chas",
        "vred_dni", "otp_dni", "otp_bud_m", "po_sredn",
        "tar1", "bal_vredn", "god", "mes"
    ]

    for r_idx, row in enumerate(df_rsv.to_dict(orient="records"), start=3):
        for c_idx, col_name in enumerate(data_cols, start=1):
            val = row.get(col_name)
            if pd.notna(val):
                ws.cell(row=r_idx, column=c_idx, value=val)

    # Итоговая формула в колонке 'Вредн' (N). Для пустой выборки диапазон
    # N3:N2 был бы некорректным, поэтому формулу пишем только при наличии строк.
    if len(df_rsv) > 0:
        last_data_row = len(df_rsv) + 2
        sum_row = last_data_row + 1
        ws.cell(row=sum_row, column=14, value=f"=SUM(N3:N{last_data_row})")

    # Автоподбор ширины колонок
    auto_fit_columns(ws)

    wb.save(out_path)


def run(
        data_dir: Path,
        month: int,
        year: int,
        stavka1: float,
        stavka2: float,
        stavka3: float,
        out_dir: Path,
        encoding: str = "cp866",
        log_callback=None,
) -> dict:
    def log(msg: str):
        if log_callback:
            log_callback(msg)
        else:
            print(msg)

    try:
        return _do_run(data_dir, month, year, stavka1, stavka2, stavka3, out_dir, encoding, log)
    except Exception as exc:
        log(f"ОШИБКА: {exc}")
        for line in traceback.format_exc().splitlines():
            log(f"    {line}")
        raise


def _do_run(
        data_dir: Path,
        month: int,
        year: int,
        stavka1: float,
        stavka2: float,
        stavka3: float,
        out_dir: Path,
        encoding: str,
        log,
) -> dict:
    m_m1 = f"{month:02d}"
    g1 = f"{year % 100:02d}"
    svod_name = f"svod{g1}{m_m1}"

    log(f"Загрузка свода: {svod_name}.dbf...")

    lschet = load_dbf(find_table(data_dir, "lschet"), encoding)
    kalend = load_dbf(find_table(data_dir, "kalend"), encoding)
    dolgn = load_dbf(find_table(data_dir, "dolgn"), encoding)
    svod = load_dbf(find_table(data_dir, svod_name), encoding)

    # 1. Фильтр записей свода по sux != 0 и vop (пустой sux считаем нулём, а не "проходит фильтр")
    df = svod[(svod["sux"].fillna(0) != 0) & (svod["vop"].isin(VOP_RSV))].copy()

    # 2. Подтяжка fio, pr_dn из lschet по tn
    lschet_clean = lschet[["tn", "fio", "pr_dn"]].drop_duplicates(subset="tn", keep="last")
    df = df.merge(lschet_clean, on="tn", how="left", suffixes=("", "_lschet"))
    if "pr_dn_lschet" in df.columns:
        df["pr_dn"] = df["pr_dn_lschet"].combine_first(df["pr_dn"])
        df.drop(columns=["pr_dn_lschet"], inplace=True)

    unmatched_lschet = df["fio"].isna()
    if unmatched_lschet.any():
        bad_tns = sorted(set(df.loc[unmatched_lschet, "tn"].tolist()))
        log(f"ВНИМАНИЕ: {unmatched_lschet.sum()} запис(ей) не сопоставлены с lschet по tn "
            f"(ФИО не найдено) — tn: {bad_tns[:20]}{' …' if len(bad_tns) > 20 else ''}")

    # 3. Подтяжка из календаря по (pr_dn + god + mes из строки свода)
    kalend_small = kalend[["pr_dn", "god", "mes", "kol_rd", "fond_ch"]].drop_duplicates(
        subset=["pr_dn", "god", "mes"]
    )
    df = df.merge(kalend_small, on=["pr_dn", "god", "mes"], how="left", suffixes=("", "_kalend"))

    # 4. Расчет sr_chas с обязательным округлением до 2 знаков (как N(5,2) в FoxPro)
    unmatched_kalend = df["kol_rd"].isna()
    if unmatched_kalend.any():
        bad_tns = sorted(set(df.loc[unmatched_kalend, "tn"].tolist()))
        log(f"ВНИМАНИЕ: {unmatched_kalend.sum()} запис(ей) не сопоставлены с календарём "
            f"(pr_dn/год/месяц) — часы/дни не рассчитаны, tn: {bad_tns[:20]}{' …' if len(bad_tns) > 20 else ''}")

    df["kol_rd"] = df["kol_rd"].fillna(0)
    df["fond_ch"] = df["fond_ch"].fillna(0)
    df["sr_chas"] = df.apply(
        lambda r: round(r["fond_ch"] / r["kol_rd"], 2) if r["kol_rd"] > 0 else 0.0, axis=1
    )

    # 4а. Защита от пустого (NaN) vf в исходных данных — не должно ронять расчёт
    missing_vf = df["vf"].isna()
    if missing_vf.any():
        bad_tns = sorted(set(df.loc[missing_vf, "tn"].tolist()))
        log(f"ВНИМАНИЕ: {missing_vf.sum()} запис(ей) с пустым vf — значение принято за 0, "
            f"tn: {bad_tns[:20]}{' …' if len(bad_tns) > 20 else ''}")
        df["vf"] = df["vf"].fillna(0)

    # 5. Раскладка дней по видам оплат
    df["vred_dni"] = 0
    df["otp_dni"] = 0
    df["otp_bud_m"] = 0
    df["po_sredn"] = 0

    is_vred = df["vop"].isin(VOP_VREDNOST)
    is_otp = df["vop"].isin(VOP_OTPUSK)
    is_otp_bud = df["vop"].isin(VOP_OTPUSK_BUD)
    is_sredn = df["vop"].isin(VOP_SREDNIY)

    zero_sr_chas = (is_vred | is_sredn) & (df["sr_chas"] <= 0)
    if zero_sr_chas.any():
        bad_tns = sorted(set(df.loc[zero_sr_chas, "tn"].tolist()))
        log(f"ВНИМАНИЕ: {zero_sr_chas.sum()} запис(ей) вредности/по средн. с sr_chas=0 "
            f"(дни не рассчитаны) — tn: {bad_tns[:20]}{' …' if len(bad_tns) > 20 else ''}")

    # Дни рассчитываются делением vf на уже округлённый sr_chas
    df.loc[is_vred, "vred_dni"] = df.loc[is_vred].apply(
        lambda r: fox_round(r["vf"] / r["sr_chas"]) if r["sr_chas"] > 0 else 0, axis=1
    )
    # fox_round(), а не astype(int): последний ОБРЕЗАЕТ дробную часть (0.5 дня
    # исчезало бы совсем), тогда как вредность и "по среднему" ниже честно
    # округляются fox_round(). Нецелый vf у отпуска — редкость, но раз он
    # встречается в данных, обрезать его молча не стоит.
    frac_otp = is_otp & (df["vf"] % 1 != 0)
    frac_otp_bud = is_otp_bud & (df["vf"] % 1 != 0)
    if (frac_otp | frac_otp_bud).any():
        bad_tns = sorted(set(df.loc[frac_otp | frac_otp_bud, "tn"].tolist()))
        log(f"ВНИМАНИЕ: {(frac_otp | frac_otp_bud).sum()} запис(ей) отпуска с дробным "
            f"числом дней (vf) — округлено, а не обрезано, tn: {bad_tns[:20]}{' …' if len(bad_tns) > 20 else ''}")
    df.loc[is_otp, "otp_dni"] = df.loc[is_otp, "vf"].apply(fox_round)
    df.loc[is_otp_bud, "otp_bud_m"] = df.loc[is_otp_bud, "vf"].apply(fox_round)
    df.loc[is_sredn, "po_sredn"] = df.loc[is_sredn].apply(
        lambda r: fox_round(r["vf"] / r["sr_chas"]) if r["sr_chas"] > 0 else 0, axis=1
    )

    # 6. Тариф и балл вредности
    df["tar1"] = df.apply(
        lambda r: round(r["sux"] / r["vf"], 2) if r["vf"] > 0 else 0.0, axis=1
    )
    df["bal_vredn"] = None
    df.loc[is_vred, "bal_vredn"] = df.loc[is_vred, "tar1"].apply(
        lambda t: classify_bal_vredn(t, stavka1, stavka2, stavka3)
    )

    unclassified = is_vred & df["bal_vredn"].isna()
    if unclassified.any():
        examples = df.loc[unclassified, ["tn", "tar1"]].head(10).values.tolist()
        log(f"ВНИМАНИЕ: {unclassified.sum()} запис(ей) вредности не удалось классифицировать по ставкам "
            f"3.1/3.2/3.3 (tar1 не совпал ни с одной из ставок) — примеры (tn, tar1): {examples}")

    # 7. Подтяжка наименования должности с обрезкой до 30 символов
    dolgn_small = dolgn[["dshifr", "dname"]].drop_duplicates(subset="dshifr")
    df = df.merge(dolgn_small, left_on="dolgn", right_on="dshifr", how="left")
    df["name"] = df["dname"].fillna("").astype(str).str[:30]
    df["empty_col"] = None

    # 8. Сортировка по no + tn
    df = df.sort_values(by=["no", "tn"]).reset_index(drop=True)

    # Сохранение RSV файла
    out_path_rsv = rsv_output_path(out_dir, year, month)
    out_path_rsv.parent.mkdir(parents=True, exist_ok=True)
    save_rsv_excel(df, out_path_rsv, month, year, stavka1, stavka2, stavka3)

    sum_vredn_rsv = int(df["vred_dni"].sum())

    log(f"Готово: {out_path_rsv} (строк: {len(df)})")
    log(f"Сумма дней вредности: {sum_vredn_rsv}")

    return {
        "rsv_path": out_path_rsv,
        "sum_rsv": sum_vredn_rsv,
        "rows_rsv": len(df),
    }


def main():
    parser = argparse.ArgumentParser(description="Расчёт вредности RSV (Clipper -> Python)")
    parser.add_argument("--data-dir", type=Path, required=True, help="Папка с .dbf файлами")
    parser.add_argument("--month", type=int, required=True, help="Месяц расчета (1-12)")
    parser.add_argument("--year", type=int, required=True, help="Год расчета, напр. 2026")
    parser.add_argument("--stavka1", type=float, default=0.29, help="Ставка 1 (по умолч. 0.29)")
    parser.add_argument("--stavka2", type=float, default=0.41, help="Ставка 2 (по умолч. 0.41)")
    parser.add_argument("--stavka3", type=float, default=0.58, help="Ставка 3 (по умолч. 0.58)")
    parser.add_argument("--out-dir", type=Path, default=Path("."), help="Папка для сохранения")
    parser.add_argument("--encoding", type=str, default="cp866", help="Кодировка DBF")
    args = parser.parse_args()

    run(
        data_dir=args.data_dir,
        month=args.month,
        year=args.year,
        stavka1=args.stavka1,
        stavka2=args.stavka2,
        stavka3=args.stavka3,
        out_dir=args.out_dir,
        encoding=args.encoding,
    )


if __name__ == "__main__":
    main()