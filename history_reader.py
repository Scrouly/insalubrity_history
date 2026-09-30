"""
history_reader.py
==================
Чтение истории по вредности из уже сформированных RSV-файлов
(insalubrity.py) — для приложения "История по сотруднику".

Работает как "обратный парсер" к save_rsv_excel() из insalubrity.py:
раскладка колонок здесь должна оставаться синхронной с той функцией.
Никакой отдельной базы данных не заводится — источник истины это
файлы out_dir/<год>/rsv_ММ.xlsx на диске.
"""

from __future__ import annotations

import json
import re
import shutil
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from app_env import app_dir, atomic_write_text
from insalubrity import VOP_VREDNOST, find_table, fox_round, load_dbf

# Порядок колонок данных в RSV должен совпадать с data_cols в save_rsv_excel()
RSV_DATA_COLUMNS = [
    "no", "tn", "fio", "vop", "vf", "sux", "dolgn", "empty_col", "name",
    "pr_dn", "fond_ch", "kol_rd", "sr_chas",
    "vred_dni", "otp_dni", "otp_bud_m", "po_sredn",
    "tar1", "bal_vredn", "god", "mes",
]

# Совпадает и новый формат (rsv_07.xlsx), и старый архивный (RSV07.xlsx,
# без подчёркивания — так называются файлы, перенесённые из старых расчётов).
RSV_FILENAME_RE = re.compile(r"^rsv_?(\d{2})\.xlsx$", re.IGNORECASE)

# От папки программы (см. app_env.app_dir), а не от текущей рабочей директории —
# иначе запуск ярлыком из другого места "терял" прежние правки (создавал новые
# пустые файлы рядом с той папкой, откуда его запустили).
DEFAULT_OVERRIDE_FILE = app_dir() / "hire_dates_override.json"
DEFAULT_MONTH_OVERRIDE_FILE = app_dir() / "month_overrides.json"

# Столбцы, которые кадровик может править вручную — только "входные" данные,
# "Итого" (monthly_total) и "Тариф" (tar1) всегда остаются расчётными.
EDITABLE_MONTH_FIELDS = ("kol_rd", "vred_dni", "bal_vredn", "otp_dni", "otp_bud_m", "po_sredn")

# Норма дополнительного отпуска за вредность (ст. 157 ТК), дней за полный
# рабочий год — по умолчанию; редактируется в интерфейсе и хранится в .env.
DEFAULT_LEAVE_NORM_DAYS = {3.1: 4, 3.2: 7, 3.3: 14}


def find_rsv_files(out_dir: Path, log=print) -> list[tuple[int, int, Path]]:
    """Находит все rsv-файлы (rsv_ММ.xlsx и архивные RSVММ.xlsx) в подпапках
    по годам: out_dir/<год>/....xlsx.

    Возвращает список (год, месяц, путь), отсортированный по (год, месяц).
    Папки, чьё имя не является 4-значным годом, и файлы, не подходящие
    под маску, молча пропускаются.

    Если за один и тот же (год, месяц) найдено несколько файлов (например,
    архивный RSV07.xlsx и новый rsv_07.xlsx лежат рядом) — берётся только
    один, с предупреждением в лог, чтобы данные не задвоились.
    """
    by_period: dict[tuple[int, int], list[Path]] = {}
    if not out_dir.exists():
        return []

    for year_dir in out_dir.iterdir():
        if not year_dir.is_dir() or not year_dir.name.isdigit():
            continue
        year = int(year_dir.name)
        for f in year_dir.glob("*.xlsx"):
            m = RSV_FILENAME_RE.match(f.name)
            if not m:
                continue
            month = int(m.group(1))
            if 1 <= month <= 12:
                by_period.setdefault((year, month), []).append(f)

    found: list[tuple[int, int, Path]] = []
    for (year, month), paths in by_period.items():
        if len(paths) > 1:
            paths.sort()
            chosen = paths[-1]
            log(
                f"ВНИМАНИЕ: за {month:02d}.{year} найдено несколько файлов "
                f"({', '.join(p.name for p in paths)}) — использован {chosen.name}, "
                f"остальные проигнорированы, чтобы избежать задвоения."
            )
        else:
            chosen = paths[0]
        found.append((year, month, chosen))

    found.sort(key=lambda t: (t[0], t[1]))
    return found


def read_rsv_file(path: Path) -> pd.DataFrame:
    """Читает строки данных (начиная с 3-й строки) из одного файла RSV.

    Останавливается на первой строке без tn (колонка B) — это либо
    строка с итоговой формулой SUM(...), либо конец данных.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = []
        for row in ws.iter_rows(min_row=3, values_only=True):
            tn = row[1] if len(row) > 1 else None
            if tn is None or (isinstance(tn, str) and not tn.strip()):
                break
            record = {
                col_name: (row[idx] if idx < len(row) else None)
                for idx, col_name in enumerate(RSV_DATA_COLUMNS)
            }
            rows.append(record)
    finally:
        wb.close()

    return pd.DataFrame(rows, columns=RSV_DATA_COLUMNS)


def build_history(out_dir: Path, log=print) -> pd.DataFrame:
    """Собирает единый DataFrame по всем найденным rsv-файлам.

    Добавляет колонки file_year/file_month, определённые из имени
    папки/файла (а не из данных внутри листа) — это самый надёжный
    источник периода, за который сделан конкретный файл.
    """
    files = find_rsv_files(out_dir, log=log)
    columns_with_period = RSV_DATA_COLUMNS + ["file_year", "file_month"]

    if not files:
        log(f"В {out_dir} не найдено ни одного файла rsv_ММ.xlsx в подпапках по годам.")
        return pd.DataFrame(columns=columns_with_period)

    frames = []
    read_periods: list[tuple[int, int]] = []
    for year, month, path in files:
        try:
            df = read_rsv_file(path)
        except Exception as exc:
            log(f"ОШИБКА чтения {path}: {exc}")
            continue
        df["file_year"] = year
        df["file_month"] = month
        frames.append(df)
        read_periods.append((year, month))

    if not frames:
        return pd.DataFrame(columns=columns_with_period)

    combined = pd.concat(frames, ignore_index=True)
    # Файл мог быть прочитан, но не дать ни одной строки (все записи отфильтрованы) —
    # тогда по одним строкам данных его существование не восстановить. Запоминаем
    # периоды явно, чтобы отличать «файла за месяц нет» от «файл есть, но у человека
    # в нём нет записей».
    combined.attrs["periods"] = sorted(read_periods)
    combined["tn"] = pd.to_numeric(combined["tn"], errors="coerce")
    for col in ("vop", "vred_dni", "otp_dni", "otp_bud_m", "po_sredn", "tar1", "kol_rd"):
        combined[col] = pd.to_numeric(combined[col], errors="coerce")
    return combined


def available_periods(history: pd.DataFrame) -> set[tuple[int, int]]:
    """Множество (год, месяц), за которые RSV-файл был прочитан."""
    periods: set[tuple[int, int]] = set()
    explicit = getattr(history, "attrs", {}).get("periods")
    if explicit:
        periods.update((int(y), int(m)) for y, m in explicit)
    if not history.empty:
        rows = history[["file_year", "file_month"]].drop_duplicates()
        periods.update((int(y), int(m)) for y, m in zip(rows["file_year"], rows["file_month"]))
    return periods


def _flag(v) -> bool:
    """Безопасное булево: NaN/None -> False (bool(nan) был бы True)."""
    return v is not None and v == v and bool(v)


def get_lschet_employees(data_dir: Path, encoding: str = "cp866") -> pd.DataFrame:
    """Список всех сотрудников (tn, fio, dnepr) из lschet.dbf — источник для автокомплита.

    Показываются все сотрудники из lschet, а не только те, у кого уже
    была вредность в rsv-файлах — иначе нового человека не найти по ФИО.
    """
    lschet = load_dbf(find_table(data_dir, "lschet"), encoding)

    required = {"tn", "fio"}
    missing = required - set(lschet.columns)
    if missing:
        raise ValueError(f"В lschet.dbf отсутствуют обязательные колонки: {missing}")

    lschet = lschet.copy()
    if "dnepr" not in lschet.columns:
        lschet["dnepr"] = None

    if "data_uvl" not in lschet.columns:
        lschet["data_uvl"] = None

    employees = (
        lschet[["tn", "fio", "dnepr", "data_uvl"]]
        .drop_duplicates(subset="tn", keep="last")
        .reset_index(drop=True)
    )
    employees["tn"] = pd.to_numeric(employees["tn"], errors="coerce")
    return employees


class OverrideFileError(Exception):
    """Файл ручных правок (даты приёма / правки месяцев) нельзя прочитать.

    Раньше в этом случае молча возвращался пустой словарь, и следующее же
    сохранение перезаписывало файл, стирая ВСЕ прежние правки. Теперь чтение
    падает громко, а запись при этом отказывается что-либо перезаписывать.
    """


MIN_HIRE_YEAR = 1950
HIRE_DATE_FUTURE_TOLERANCE_DAYS = 31
# Плановое увольнение (дата внесена заранее) — планы на 3 года вперёд уже
# подозрительны и скорее говорят об ошибке ввода, чем о реальном кадровом плане.
TERMINATION_DATE_FUTURE_TOLERANCE_DAYS = 366 * 3


def validate_termination_date(d) -> date:
    """Проверяет правдоподобие даты увольнения (DATA_UVL). Возвращает date
    или бросает ValueError. Планового увольнения в далёком будущем (годы за
    сотни лет вперёд — обычно признак сбоя формата даты в DBF) не бывает."""
    if isinstance(d, pd.Timestamp):
        d = d.date()
    if not isinstance(d, date):
        raise ValueError(f"Дата увольнения не распознана: {d!r}")
    if d.year < MIN_HIRE_YEAR:
        raise ValueError(f"Дата увольнения {d.isoformat()} слишком ранняя (раньше {MIN_HIRE_YEAR} г.)")
    if d > date.today() + timedelta(days=TERMINATION_DATE_FUTURE_TOLERANCE_DAYS):
        raise ValueError(f"Дата увольнения {d.isoformat()} нереалистично далеко в будущем")
    return d


def validate_hire_date(d) -> date:
    """Проверяет правдоподобие даты приёма. Возвращает date или бросает ValueError."""
    if isinstance(d, pd.Timestamp):
        d = d.date()
    if not isinstance(d, date):
        raise ValueError(f"Дата приёма не распознана: {d!r}")
    if d.year < MIN_HIRE_YEAR:
        raise ValueError(f"Дата приёма {d.isoformat()} слишком ранняя (раньше {MIN_HIRE_YEAR} г.)")
    if d > date.today() + timedelta(days=HIRE_DATE_FUTURE_TOLERANCE_DAYS):
        raise ValueError(f"Дата приёма {d.isoformat()} находится в будущем")
    return d


def _load_json_dict(path: Path) -> dict:
    path = Path(path)
    if not path.exists():
        return {}
    try:
        text = path.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        raise OverrideFileError(f"Не удалось прочитать файл {path}: {exc}") from exc
    if not text.strip():
        return {}  # пустой файл — то же, что «правок ещё нет»
    try:
        data = json.loads(text)
    except json.JSONDecodeError as exc:
        raise OverrideFileError(
            f"Файл {path} повреждён ({exc}).\n"
            f"Он НЕ был изменён. Исправьте его вручную или восстановите из "
            f"{path.name}.bak (если есть), иначе новые правки сохранять нельзя."
        ) from exc
    if not isinstance(data, dict):
        raise OverrideFileError(f"Файл {path} имеет неожиданную структуру (ожидался объект JSON).")
    return data


def _save_json_dict(path: Path, data: dict) -> None:
    """Атомарная запись + одна резервная копия предыдущей версии (.bak)."""
    path = Path(path)
    if path.exists():
        try:
            shutil.copy2(path, path.with_name(path.name + ".bak"))
        except OSError:
            pass  # отсутствие бэкапа не должно блокировать сохранение
    atomic_write_text(path, json.dumps(data, ensure_ascii=False, indent=2))


def load_hire_date_overrides(path: Path = DEFAULT_OVERRIDE_FILE) -> dict[str, str]:
    return _load_json_dict(path)


def save_hire_date_override(tn: int, date_str: str, path: Path = DEFAULT_OVERRIDE_FILE) -> None:
    try:
        parsed = pd.to_datetime(date_str)
    except Exception as exc:
        raise ValueError(f"Дата приёма не распознана: {date_str!r}") from exc
    if pd.isna(parsed):
        raise ValueError(f"Дата приёма не распознана: {date_str!r}")
    validate_hire_date(parsed.date())

    overrides = load_hire_date_overrides(path)  # при битом файле бросит OverrideFileError
    overrides[str(int(tn))] = parsed.date().isoformat()
    _save_json_dict(path, overrides)


def load_month_overrides(path: Path = DEFAULT_MONTH_OVERRIDE_FILE) -> dict:
    """Возвращает {tn_str: {"YYYY-MM": {field: value, ...}, ...}, ...}."""
    return _load_json_dict(path)


# Служебный ключ внутри записи месяца: были ли в RSV НАСТОЯЩИЕ данные за этот
# месяц в момент, когда кадровик вводил правку. Если данных не было (правка
# «на будущее»/поверх оценки), а потом появился реальный RSV с другим
# значением — правка помечается устаревшей (см. apply_month_overrides).
HAD_DATA_KEY = "_had_data"


def save_month_override(
    tn: int, year: int, month: int, field: str, value,
    path: Path = DEFAULT_MONTH_OVERRIDE_FILE,
    had_data: bool | None = None,
) -> None:
    if field not in EDITABLE_MONTH_FIELDS:
        raise ValueError(f"Поле '{field}' нельзя редактировать вручную")
    overrides = load_month_overrides(path)
    tn_key = str(int(tn))
    ym_key = f"{year:04d}-{month:02d}"
    entry = overrides.setdefault(tn_key, {}).setdefault(ym_key, {})
    entry[field] = value
    if had_data is not None:
        # Повторный ввод значения = подтверждение правки при текущем состоянии данных.
        entry[HAD_DATA_KEY] = bool(had_data)
    _save_json_dict(path, overrides)


def clear_month_override(
    tn: int, year: int, month: int, field: str,
    path: Path = DEFAULT_MONTH_OVERRIDE_FILE,
) -> None:
    """Убирает ручную правку одного поля — значение вернётся к тому, что даёт
    RSV (если за этот месяц есть настоящие данные) или прогноз/ноль (если нет)."""
    overrides = load_month_overrides(path)
    tn_key = str(int(tn))
    ym_key = f"{year:04d}-{month:02d}"
    if tn_key in overrides and ym_key in overrides[tn_key]:
        overrides[tn_key][ym_key].pop(field, None)
        # Если кроме служебного маркера ничего не осталось — запись месяца удаляем целиком.
        if not any(k in overrides[tn_key][ym_key] for k in EDITABLE_MONTH_FIELDS):
            del overrides[tn_key][ym_key]
        if not overrides[tn_key]:
            del overrides[tn_key]
        _save_json_dict(path, overrides)


def get_hire_date(tn: int, employees: pd.DataFrame, overrides: dict[str, str] | None = None):
    """Возвращает (date_или_None, источник).

    Источник: 'override' / 'dnepr' / 'dnepr_invalid' / None.
    Приоритет: ручная правка кадровика (override) > поле DNEPR из lschet.dbf.
    'dnepr_invalid' — в DNEPR что-то есть, но дата неправдоподобна (например,
    1899 г. или будущее) — такая дата не используется, чтобы не строить
    таймлайн на сотни месяцев от «нулевой» даты.
    """
    if overrides is None:
        overrides = load_hire_date_overrides()

    key = str(int(tn))
    if overrides.get(key):
        try:
            parsed = pd.to_datetime(overrides[key])
            if pd.notna(parsed):
                return validate_hire_date(parsed.date()), "override"
        except Exception:
            pass  # битая ручная дата — падаем обратно на DNEPR

    row = employees.loc[employees["tn"] == tn]
    if not row.empty:
        raw = row.iloc[0].get("dnepr")
        if raw not in (None, "", 0, "0"):
            try:
                parsed = pd.to_datetime(raw)
            except Exception:
                return None, "dnepr_invalid"
            if pd.notna(parsed):
                try:
                    return validate_hire_date(parsed.date()), "dnepr"
                except ValueError:
                    return None, "dnepr_invalid"

    return None, None


def get_termination_date(tn: int, employees: pd.DataFrame, hire_date=None):
    """Возвращает (date_или_None, источник) для DATA_UVL из lschet.dbf.

    Источник: 'data_uvl' / 'data_uvl_invalid' (неправдоподобная дата) /
    'data_uvl_before_hire' (увольнение раньше даты приёма — типичный случай
    переприёма: старая DATA_UVL из предыдущего периода работы, сейчас
    сотрудник снова числится с более поздней DNEPR/override'ом — такую
    дату увольнения игнорируем, иначе таймлайн обрывался бы посреди текущей
    работы) / None (сотрудник действующий, DATA_UVL пусто).
    """
    row = employees.loc[employees["tn"] == tn]
    if row.empty:
        return None, None
    raw = row.iloc[0].get("data_uvl")
    if raw in (None, "", 0, "0"):
        return None, None

    try:
        parsed = pd.to_datetime(raw)
    except Exception:
        return None, "data_uvl_invalid"
    if pd.isna(parsed):
        return None, None

    try:
        term_date = validate_termination_date(parsed.date())
    except ValueError:
        return None, "data_uvl_invalid"

    if hire_date is not None and term_date < hire_date:
        return None, "data_uvl_before_hire"

    return term_date, "data_uvl"


def find_data_before_hire(tn: int, history: pd.DataFrame, hire_date) -> list[tuple[int, int]]:
    """Месяцы (год, месяц), за которые у сотрудника есть RSV-данные РАНЬШЕ
    месяца приёма. build_employee_timeline начинает таймлайн с даты приёма,
    поэтому такие месяцы в таблицу не попадают — GUI должен предупредить об
    этом (типичная причина: повторный приём, DNEPR обновлён новой датой)."""
    if hire_date is None or history is None or history.empty:
        return []
    person = history.loc[history["tn"] == tn, ["file_year", "file_month"]].drop_duplicates()
    hire_key = (hire_date.year, hire_date.month)
    early = sorted(
        (int(y), int(m)) for y, m in zip(person["file_year"], person["file_month"])
        if (int(y), int(m)) < hire_key
    )
    return early


def _nz(v):
    """None/NaN -> 0, иначе само значение. `x or 0` для этого НЕ годится:
    NaN истинен, поэтому `nan or 0` остаётся nan и ломает round()/int()."""
    return 0 if v is None or pd.isna(v) else v


def work_year_bounds(hire_date, year: int, month: int) -> tuple[int, int, int, int]:
    """Возвращает (start_year, start_month, end_year, end_month) рабочего года
    (ст. 163 ТК — 12 месяцев со дня приёма на работу, а не календарный год),
    в который попадает календарный (year, month).

    Упрощение: рабочий год считается фиксированными 12-месячными блоками от
    даты приёма, без сдвига по ст. 165 ТК (сдвиг требует данных о днях,
    исключаемых из рабочего года — отпуска за свой счёт свыше 14 дней и т.п.,
    которых нет в RSV/своде начислений).
    """
    hire_idx = hire_date.year * 12 + (hire_date.month - 1)
    cur_idx = year * 12 + (month - 1)
    n = (cur_idx - hire_idx) // 12
    start_idx = hire_idx + n * 12
    end_idx = start_idx + 11
    start_year, start_month = divmod(start_idx, 12)
    start_month += 1
    end_year, end_month = divmod(end_idx, 12)
    end_month += 1
    return start_year, start_month, end_year, end_month


def _apply_bal_vredn_ffill(result: pd.DataFrame) -> pd.DataFrame:
    """Переносит последний известный класс вредности (bal_vredn) вперёд по
    хронологии — в bal_vredn_effective — и попутно помечает, каким месяцам
    класс достался НЕ из собственной RSV-строки, а перенесён с более раннего
    месяца: bal_vredn_carried=True, плюс bal_vredn_source_year/month — год и
    месяц, ОТКУДА он перенесён (для тултипа в GUI). Учитывает, что
    result уже отсортирован по (year, month) по построению.

    Вызывается и при первичной сборке таймлайна, и заново — после applying
    ручных правок (apply_month_overrides), т.к. правка bal_vredn на одном
    месяце меняет перенос для всех последующих месяцев.
    """
    raw = result["bal_vredn"]
    result["bal_vredn_effective"] = raw.ffill()
    result["bal_vredn_carried"] = raw.isna() & result["bal_vredn_effective"].notna()
    result["bal_vredn_source_year"] = result["year"].where(raw.notna()).ffill()
    result["bal_vredn_source_month"] = result["month"].where(raw.notna()).ffill()
    return result


def build_employee_timeline(
    tn: int,
    history: pd.DataFrame,
    hire_date=None,
    termination_date=None,
) -> pd.DataFrame:
    """Строит помесячный таймлайн вредности для сотрудника.

    Возвращает DataFrame с колонками:
    год, месяц, vred_dni, bal_vredn, otp_dni, otp_bud_m, po_sredn, tar1, has_data
    от месяца устройства (или самого раннего найденного месяца по этому
    tn, если дата устройства неизвестна) до последнего месяца, за который
    вообще есть данные в out_dir (по всем сотрудникам — чтобы месяцы без
    расчёта не выглядели как "сотрудник не найден").

    Месяцы без записей по сотруднику включаются в таймлайн с нулями,
    а не пропускаются.

    termination_date (DATA_UVL из lschet.dbf, см. get_termination_date):
    если задана, таймлайн НЕ строится дальше месяца увольнения — иначе
    последующие "месяцев без данных" были бы позже автоматически заменены
    оценочным средним (fix_zero_kol_rd_months) и/или достроены на весь
    рабочий год вперёд (project_current_work_year), приписывая уволенному
    сотруднику вредность и норму дней за время, когда он уже не работал.
    Настоящие RSV-записи после даты увольнения (например, финальный расчёт
    следующим месяцем) НИКОГДА не отбрасываются — граница отодвигается до
    последнего месяца с реальными данными, если такой позже увольнения.
    """
    empty_columns = [
        "year", "month", "kol_rd", "vred_dni", "bal_vredn",
        "otp_dni", "otp_bud_m", "po_sredn", "tar1", "has_data",
        "work_year_key", "work_year_label", "monthly_total", "bal_vredn_effective",
        "is_projected",
        "is_manual_override",
        "is_data_fix",
        "no_file",
    ]

    if history.empty:
        return pd.DataFrame(columns=empty_columns)

    person = history.loc[history["tn"] == tn].copy()

    # Последний месяц, за который вообще есть данные (по всем сотрудникам)
    periods = available_periods(history)
    latest_year, latest_month = max(periods)

    if termination_date is not None:
        end_key = (termination_date.year, termination_date.month)
        if not person.empty:
            last_real_key = (
                int(person["file_year"].max()),
                int(person.loc[person["file_year"] == person["file_year"].max(), "file_month"].max()),
            )
            if last_real_key > end_key:
                end_key = last_real_key  # реальные данные важнее формальной даты увольнения
        if end_key < (latest_year, latest_month):
            latest_year, latest_month = end_key

    if hire_date is not None:
        start_year, start_month = hire_date.year, hire_date.month
    elif not person.empty:
        start_year = int(person["file_year"].min())
        start_month = int(person.loc[person["file_year"] == start_year, "file_month"].min())
    else:
        # Ни даты устройства, ни единой записи по этому tn — таймлайн пуст
        return pd.DataFrame(columns=empty_columns)

    # Агрегация по месяцу (на случай нескольких строк с разными vop за один месяц)
    if not person.empty:
        grouped = person.groupby(["file_year", "file_month"], as_index=False).agg(
            # kol_rd — норма рабочих дней месяца по графику (константа для
            # сотрудника в этом месяце), а не сумма по строкам вредности/отпуска.
            kol_rd=("kol_rd", "max"),
            vred_dni=("vred_dni", "sum"),
            otp_dni=("otp_dni", "sum"),
            otp_bud_m=("otp_bud_m", "sum"),
            po_sredn=("po_sredn", "sum"),
            bal_vredn=("bal_vredn", "max"),
        )
        grouped = grouped.set_index(["file_year", "file_month"])

        # Тариф (tar1) берём ТОЛЬКО из строк вредности (vop=49). У строк
        # отпуска/по-среднему tar1 = sux/vf — это рубли за день, а не ставка
        # вредности, и max() по месяцу подсовывал в колонку «Тариф» именно его.
        # Строки с неизвестным vop (старые архивные файлы) оставляем как раньше.
        if "vop" in person.columns:
            vop_num = pd.to_numeric(person["vop"], errors="coerce")
            tar_src = person.loc[(vop_num == min(VOP_VREDNOST)) | vop_num.isna()]
        else:
            tar_src = person
        tar1_by_month = tar_src.groupby(["file_year", "file_month"])["tar1"].max()
        grouped["tar1"] = tar1_by_month.reindex(grouped.index)
    else:
        grouped = pd.DataFrame(
            columns=["kol_rd", "vred_dni", "otp_dni", "otp_bud_m", "po_sredn", "tar1", "bal_vredn"]
        )

    rows = []
    y, m = start_year, start_month
    while (y, m) <= (latest_year, latest_month):
        if hire_date is not None:
            wy_start_y, wy_start_m, wy_end_y, wy_end_m = work_year_bounds(hire_date, y, m)
            work_year_key = (wy_start_y, wy_start_m)
            work_year_label = f"{wy_start_m:02d}.{wy_start_y}–{wy_end_m:02d}.{wy_end_y}"
        else:
            # Дата приёма неизвестна — рабочий год посчитать нельзя,
            # используем календарный как приближение (с пометкой в label).
            work_year_key = (y,)
            work_year_label = f"{y} (календарный — дата приёма неизвестна)"

        if (y, m) in grouped.index:
            rec = grouped.loc[(y, m)]
            v, o, ob, ps = _nz(rec["vred_dni"]), _nz(rec["otp_dni"]), _nz(rec["otp_bud_m"]), _nz(rec["po_sredn"])
            rows.append({
                "year": y, "month": m,
                "kol_rd": _nz(rec["kol_rd"]),
                "vred_dni": v,
                "bal_vredn": rec["bal_vredn"],
                "otp_dni": o,
                "otp_bud_m": ob,
                "po_sredn": ps,
                "tar1": rec["tar1"],
                "has_data": True,
                "work_year_key": work_year_key,
                "work_year_label": work_year_label,
                "monthly_total": v + o + ob + ps,
                "is_projected": False,
                "is_manual_override": False,
                "is_data_fix": False,
                "no_file": False,
            })
        else:
            rows.append({
                "year": y, "month": m,
                "kol_rd": 0,
                "vred_dni": 0, "bal_vredn": None,
                "otp_dni": 0, "otp_bud_m": 0, "po_sredn": 0, "tar1": None,
                "has_data": False,
                "work_year_key": work_year_key,
                "work_year_label": work_year_label,
                "monthly_total": 0,
                "is_projected": False,
                "is_manual_override": False,
                "is_data_fix": False,
                # True — RSV-файла за этот месяц нет вообще (не рассчитан/не прочитан/
                # архив начинается позже). False при has_data=False — файл есть, но у
                # сотрудника в нём нет записей.
                "no_file": (y, m) not in periods,
            })
        m += 1
        if m > 12:
            m = 1
            y += 1

    result = pd.DataFrame(rows, columns=empty_columns)

    # Починка сбоя календаря: RSV-строка за месяц ЕСТЬ (человек точно
    # работал), но норма рабочих дней (kol_rd) ошибочно 0 — это дефект
    # сопоставления с kalend.dbf в исходном расчёте, а не факт "0 дней
    # по графику". Делаем это ДО переноса класса вредности и ДО
    # прогноза текущего года — почин ённые месяцы затем корректно
    # участвуют в пуле усреднения для project_current_work_year.
    result = fix_zero_kol_rd_months(result)

    # Перенос последнего известного класса вредности вперёд: если человек
    # работает во вредных условиях, то отпуск/о-б-м/по-среднему в том же
    # рабочем месте продолжают относиться к тому же классу, даже если в этом
    # конкретном месяце не было отдельной строки вредности (vop=49). Перенос
    # идёт по всей хронологии сотрудника (не сбрасывается на границе рабочего
    # года), т.к. должность не меняется сама по себе от смены раб. года.
    # Остаются неклассифицированными только месяцы ДО первого когда-либо
    # встреченного класса (когда переносить попросту нечего).
    result = _apply_bal_vredn_ffill(result)

    return result


def fix_zero_kol_rd_months(timeline: pd.DataFrame) -> pd.DataFrame:
    """Чинит ЛЮБОЙ месяц, где норма рабочих дней по графику (kol_rd) равна 0,
    хотя должна быть известна — двух видов, с разной причиной и разной
    пометкой в GUI:

    1. has_data=True (в RSV за этот месяц ЕСТЬ запись — человек точно
       числился), но kol_rd=0 — это дефект сопоставления с календарём
       (kalend.dbf) в исходном расчёте (insalubrity.py: kol_rd/fond_ch
       подтягиваются по (pr_dn, год, месяц), и при несовпадении остаются
       нулевыми). Помечается is_data_fix=True — в GUI бирюзовый акцент
       "починено".

    2. has_data=False (в RSV за этот месяц вообще НЕТ ни одной записи —
       insalubrity.py целиком выбрасывает такие месяцы фильтром sux!=0 ещё
       до сопоставления с календарём, например если человек этот месяц не
       был занят во вредных условиях). Норма дней по графику для такого
       месяца оценивается тем же способом. Помечается is_projected=True —
       в GUI это тот же зелёный акцент "(оценка)", что и у прогноза на
       будущее, т.к. по сути это то же самое: настоящих данных нет, значение
       расчётное.

    В обоих случаях kol_rd заменяется на среднее по последним (до 12,
    хронологически) месяцам ЭТОГО ЖЕ сотрудника из ВСЕЙ его истории, где
    kol_rd > 0 — или по всем таким месяцам, если их меньше 12 (та же логика
    источника, что и в project_current_work_year). Если во всей истории
    сотрудника вообще нет ни одного месяца с ненулевым kol_rd — чинить
    нечем, значение остаётся 0.

    vred_dni (дни вредности) сознательно НЕ трогается ни в одном из двух
    случаев и остаётся тем, что было (обычно тоже 0) — подменять фактически
    начисленные/отсутствующие дни вредности расчётным средним не нужно,
    оценивается только норма дней по графику. Именно поэтому has_data НЕ
    меняется на True в случае 2 — по вредности данных как не было, так и
    нет, только теперь известна норма дней по графику для точного расчёта
    доли рабочего года.
    """
    if timeline.empty:
        return timeline

    result = timeline.copy()
    if "data_fix_note" not in result.columns:
        result["data_fix_note"] = None
    if "projection_note" not in result.columns:
        result["projection_note"] = None
    if "no_file" not in result.columns:
        result["no_file"] = False

    broken_mask = result["kol_rd"].fillna(0) == 0
    if not broken_mask.any():
        return result

    # kol_rd > 0 уже само по себе означает "настоящий известный график" —
    # has_data=False месяцы здесь никогда не попадут (у них kol_rd всегда 0
    # по построению в build_employee_timeline), доп. фильтр не нужен.
    pool = result[result["kol_rd"] > 0].sort_values(["year", "month"]).tail(12)
    if pool.empty:
        return result  # во всей истории сотрудника нет ни одного месяца-образца — чинить нечем

    avg_kol_rd = fox_round(float(pool["kol_rd"].mean()))
    pool_size = len(pool)

    has_data_broken = broken_mask & result["has_data"]
    no_data_broken = broken_mask & ~result["has_data"]

    if has_data_broken.any():
        fix_note = (
            f"Раб. дни: в RSV было 0 (сбой сопоставления с календарём), "
            f"авто-заменено на среднее по {pool_size} последним отработанным мес.: {avg_kol_rd}"
        )
        result.loc[has_data_broken, "kol_rd"] = avg_kol_rd
        result.loc[has_data_broken, "is_data_fix"] = True
        result.loc[has_data_broken, "data_fix_note"] = fix_note

    if no_data_broken.any():
        proj_note = (
            f"Оценка: за этот месяц в RSV нет ни одной записи (не было начислений по "
            f"вредности/отпуску/по-среднему) — норма дней по графику оценена как среднее "
            f"по {pool_size} последним отработанным мес.: {avg_kol_rd}"
        )
        result.loc[no_data_broken, "kol_rd"] = avg_kol_rd
        result.loc[no_data_broken, "is_projected"] = True
        result.loc[no_data_broken, "projection_note"] = proj_note
        no_file_broken = no_data_broken & result["no_file"].map(_flag)
        if no_file_broken.any():
            result.loc[no_file_broken, "projection_note"] = (
                f"RSV-файла за этот месяц нет (не рассчитан или не прочитан) — реальных "
                f"данных нет; норма дней по графику оценена как среднее по {pool_size} "
                f"последним отработанным мес.: {avg_kol_rd}"
            )

    return result


def summarize_leave_entitlement(
    timeline: pd.DataFrame,
    norm_days: dict[float, float] | None = None,
) -> list[dict]:
    """Считает дополнительный отпуск за вредность по каждому рабочему году.

    Логика (см. ст. 157 ТК): доп. отпуск за вредность считается пропорционально
    фактически отработанному времени во вредных условиях за рабочий год:

        вышло = Σ по классам (дни этого класса за раб. год × норма класса) / норма
                рабочих дней за раб. год по графику

    "Дни этого класса" — это monthly_total (вредность + отпуск + о/б/м + по
    среднему) месяцев, отнесённых к данному классу через "bal_vredn_effective"
    (класс, перенесённый вперёд с последнего месяца, где он был явно определён —
    см. build_employee_timeline). Пока человек работает во вредных условиях,
    отпуск и прочие выплаты по этой же должности считаются тем же классом, даже
    если в конкретном месяце не было отдельной строки вредности.

    Неклассифицированными (unclassified_days) остаются только месяцы ДО первого
    когда-либо определённого класса у сотрудника — переносить ещё нечего.

    Рабочие года, где вредность (vred_dni) за ВЕСЬ год равна 0 (человек за это
    время реально ни дня не отработал во вредных условиях — даже если старый
    класс формально всё ещё "тянется" по ffill и в этом году были отпуск/по-
    среднему), в результат вообще не попадают: отпуск за вредность для них не
    считается.

    "положено" — сумма норм каждого встретившегося за раб. год класса (без
    пропорционального пересчёта) — справочное значение, сколько причиталось
    бы за полностью отработанный рабочий год в этом классе(-ах).

    Возвращает список словарей (по одному на рабочий год, в хронологическом
    порядке), каждый с ключами:
        work_year_key, work_year_label, kol_rd_total,
        classes: [{"klass": 3.1, "days": ..., "norm": ..., "vyshlo": ...}, ...],
        unclassified_days, polozheno_total, vyshlo_total
    """
    if norm_days is None:
        norm_days = DEFAULT_LEAVE_NORM_DAYS

    if timeline.empty:
        return []

    results = []
    for work_year_key, group in timeline.groupby("work_year_key", sort=False):
        work_year_label = group["work_year_label"].iloc[0]
        kol_rd_total = float(group["kol_rd"].fillna(0).sum())

        # Если за ВЕСЬ рабочий год не было ни одного дня, реально отработанного
        # во вредных условиях (vred_dni), отпуск за вредность для этого года не
        # считаем вообще — даже если бывший класс всё ещё "тянется" по ffill
        # с давнего месяца и в этом году набежали дни отпуска/по-среднему.
        # Без этой проверки человеку, который уже год как не работает во
        # вредных условиях (а просто был в обычном отпуске под старым
        # классом), продолжал бы начисляться доп. отпуск на пустом месте.
        vred_total = float(group["vred_dni"].fillna(0).sum())
        if vred_total <= 0:
            continue

        # Округляем класс до 1 знака, чтобы избежать сюрпризов с плавающей точкой.
        # ВАЖНО: не строить это через .apply(lambda: ... else None) — pandas
        # приводит результат обратно к float64, и None превращается в NaN;
        # а NaN != NaN, из-за чего в set/dict каждый NaN попадает как
        # отдельный "уникальный" ключ. Используем .round() напрямую — NaN
        # остаётся NaN по всей колонке, и isna()/dropna() работают корректно.
        # to_numeric нужен отдельно: если класс НИ РАЗУ не был определён за
        # всю историю сотрудника, вся колонка — object/None (а не float/NaN),
        # и .round() на такой колонке падает с TypeError.
        rounded = pd.to_numeric(group["bal_vredn_effective"], errors="coerce").round(1)

        class_totals: dict[float, float] = {}
        for c in sorted(rounded.dropna().unique()):
            class_totals[c] = float(group.loc[rounded == c, "monthly_total"].fillna(0).sum())

        unclassified_days = float(group.loc[rounded.isna(), "monthly_total"].fillna(0).sum())

        classes_breakdown = []
        polozheno_total = 0.0
        vyshlo_total = 0.0
        for c in sorted(class_totals.keys()):
            days = class_totals[c]
            norm = norm_days.get(c, 0)
            vyshlo_raw = (days * norm / kol_rd_total) if kol_rd_total > 0 else 0.0
            # Больше нормы за полный рабочий год по классу не положено — иначе
            # избыточные дни (например, отпуск в календарных днях поверх рабочего
            # графика) раздували бы результат выше законной нормы.
            vyshlo = min(vyshlo_raw, norm)
            capped = vyshlo_raw > norm + 1e-9
            classes_breakdown.append({
                "klass": c, "days": days, "norm": norm,
                "vyshlo": vyshlo, "vyshlo_raw": vyshlo_raw, "capped": capped,
            })
            polozheno_total += norm
            vyshlo_total += vyshlo

        # Сумма «дней» больше нормы рабочих дней графика — верный признак смешения
        # единиц (отпуск в календарных днях) или неверной нормы kol_rd.
        total_days = sum(class_totals.values()) + unclassified_days
        days_exceed_norm = kol_rd_total > 0 and total_days > kol_rd_total + 1e-9

        results.append({
            "work_year_key": work_year_key,
            "work_year_label": work_year_label,
            "kol_rd_total": kol_rd_total,
            "classes": classes_breakdown,
            "unclassified_days": unclassified_days,
            "days_exceed_norm": days_exceed_norm,
            "total_days": total_days,
            "polozheno_total": polozheno_total,
            "vyshlo_total": vyshlo_total,
        })

    return results


def project_current_work_year(timeline: pd.DataFrame, termination_date=None) -> pd.DataFrame:
    """Достраивает ТЕКУЩИЙ (последний) рабочий год до полных 12 месяцев
    прогнозными значениями kol_rd/vred_dni там, где реальных RSV-данных ещё
    нет — иначе расчёт отпуска занижает годовой итог только из-за того, что
    год ещё не закончился (RSV обычно отстаёт от календаря на 1-2 месяца).

    termination_date: если задана, достройка не заходит дальше месяца
    увольнения — сотрудник, которого уже нет, не может "доработать" год
    вперёд. build_employee_timeline уже не строит исходные строки позже
    этой даты, так что без этой проверки достройка сама подставила бы
    туда прогнозные месяцы задним числом.

    Источник прогноза (берётся первый непустой вариант):
      1. Предыдущий рабочий год (блок прямо перед текущим) — среднее по его
         месяцам, где ОБА показателя (kol_rd И vred_dni) одновременно не 0.
      2. Последние 12 фактических месяцев из всей истории сотрудника — по
         тому же правилу "оба показателя не 0".
      3. Все фактические месяцы истории — по тому же правилу.
    Если ни один вариант не даёт ни одного подходящего месяца — прогноз не
    строится, таймлайн возвращается как есть.

    Спрогнозированные месяцы получают has_data=False (это не факт) и
    is_projected=True — GUI показывает их с пометкой «(оценка)». Класс
    вредности берётся как последний известный (bal_vredn_effective) —
    это атрибут должности, его не усредняют.
    """
    if timeline.empty:
        return timeline

    result = timeline.copy()

    keys_in_order: list[tuple] = []
    seen = set()
    for k in result["work_year_key"]:
        if k not in seen:
            seen.add(k)
            keys_in_order.append(k)

    last_key = keys_in_order[-1]
    current_block = result[result["work_year_key"] == last_key]
    if len(current_block) >= 12:
        return result  # рабочий год уже полный — достраивать нечего

    if len(last_key) == 2:
        # Обычный рабочий год (дата устройства известна): 12 месяцев от даты устройства.
        start_year, start_month = last_key
    else:
        # Дата устройства неизвестна — группировка идёт по календарному году
        # (см. build_employee_timeline), поэтому "рабочий год" здесь — просто
        # календарь янв–дек того же года.
        start_year, start_month = last_key[0], 1
    start_idx = start_year * 12 + (start_month - 1)
    all_months = []
    for i in range(12):
        idx = start_idx + i
        y, m = divmod(idx, 12)
        all_months.append((y, m + 1))

    existing_months = set(zip(current_block["year"].astype(int), current_block["month"].astype(int)))
    # Достраиваем ТОЛЬКО вперёд, после самого позднего уже известного месяца
    # (реального или зафиксированного placeholder'ом) — иначе при неизвестной
    # дате устройства (work_year_key = календарный год, начинается с января)
    # сюда попадут месяцы ДО того, как по сотруднику вообще появилась первая
    # запись, что выглядит как вредность до трудоустройства.
    max_existing = max(existing_months)
    missing_months = [ym for ym in all_months if ym not in existing_months and ym > max_existing]
    if termination_date is not None:
        term_key = (termination_date.year, termination_date.month)
        missing_months = [ym for ym in missing_months if ym <= term_key]
    if not missing_months:
        return result

    def qualifying(df: pd.DataFrame) -> pd.DataFrame:
        return df[df["has_data"] & (df["kol_rd"] > 0) & (df["vred_dni"] > 0)]

    pool = pd.DataFrame()
    source_note = ""
    if len(keys_in_order) >= 2:
        prev_block = result[result["work_year_key"] == keys_in_order[-2]]
        pool = qualifying(prev_block)
        if not pool.empty:
            source_note = f"среднее по предыдущему рабочему году ({len(pool)} мес. с данными)"

    if pool.empty:
        actual_sorted = result[result["has_data"]].sort_values(["year", "month"])
        pool = qualifying(actual_sorted.tail(12))
        if not pool.empty:
            source_note = f"среднее за последние {len(pool)} фактических мес."

    if pool.empty:
        pool = qualifying(result[result["has_data"]])
        if not pool.empty:
            source_note = f"среднее за весь имеющийся период ({len(pool)} мес. с данными)"

    if pool.empty:
        return result  # не из чего строить прогноз

    avg_kol_rd = fox_round(float(pool["kol_rd"].mean()))
    avg_vred_dni = fox_round(float(pool["vred_dni"].mean()))

    known_rows = result.loc[result["bal_vredn_effective"].notna()]
    if not known_rows.empty:
        last_known = known_rows.iloc[-1]
        projected_class = last_known["bal_vredn_effective"]
        # Источник переноса — тот же самый месяц, что был найден для
        # last_known (даже если last_known сам уже был перенесённым, а не
        # исходной RSV-строкой), чтобы тултип всегда указывал на настоящий
        # первоисточник класса, а не на промежуточное звено цепочки.
        src_year = last_known["bal_vredn_source_year"]
        src_month = last_known["bal_vredn_source_month"]
        if pd.isna(src_year) or pd.isna(src_month):
            src_year, src_month = last_known["year"], last_known["month"]
    else:
        projected_class = None
        src_year, src_month = None, None

    label = current_block.iloc[0]["work_year_label"]
    new_rows = [
        {
            "year": y, "month": m,
            "kol_rd": avg_kol_rd,
            "vred_dni": avg_vred_dni,
            "bal_vredn": None,
            "otp_dni": 0, "otp_bud_m": 0, "po_sredn": 0, "tar1": None,
            "has_data": False,
            "work_year_key": last_key,
            "work_year_label": label,
            "monthly_total": avg_vred_dni,
            "bal_vredn_effective": projected_class,
            # Спрогнозированный месяц по определению не подтверждён RSV, так
            # что класс здесь ВСЕГДА перенесённый (если вообще есть откуда
            # переносить) — GUI красит "Балл" фиолетовым так же, как и в
            # обычных (не прогнозных) перенесённых месяцах.
            "bal_vredn_carried": projected_class is not None,
            "bal_vredn_source_year": src_year,
            "bal_vredn_source_month": src_month,
            "is_projected": True,
            "is_manual_override": False,
            "is_data_fix": False,
            "no_file": False,
            "data_fix_note": None,
            "projection_note": source_note,
        }
        for y, m in missing_months
    ]

    result = pd.concat([result, pd.DataFrame(new_rows)], ignore_index=True)
    return result.sort_values(["year", "month"]).reset_index(drop=True)


def _values_equal(a, b) -> bool:
    a_na, b_na = pd.isna(a), pd.isna(b)
    if a_na or b_na:
        return bool(a_na and b_na)
    try:
        return abs(float(a) - float(b)) < 1e-6
    except (TypeError, ValueError):
        return a == b


def apply_month_overrides(timeline: pd.DataFrame, tn_overrides: dict) -> pd.DataFrame:
    """Накладывает ручные правки кадровика на месяцы — ЛЮБЫЕ, включая те, где
    уже есть настоящие RSV-данные. Ручная правка теперь всегда в приоритете
    над RSV/прогнозом: если кадровик поправил месяц, это значение и
    используется, пока правку явно не уберут (clear_month_override) —
    независимо от того, что впоследствии появится в RSV за этот месяц.

    Для каждой переопределённой ЯЧЕЙКИ (а не строки целиком) в колонке
    "overridden_fields" запоминается исходное значение (то, что было в
    RSV/прогнозе до правки) — GUI показывает его во всплывающей подсказке и
    подсвечивает только эту ячейку, а не всю строку.

    tn_overrides — под-словарь для ОДНОГО сотрудника из load_month_overrides(),
    вида {"YYYY-MM": {"kol_rd": .., "vred_dni": .., "bal_vredn": .., ...}}.
    """
    if not tn_overrides or timeline.empty:
        return timeline

    result = timeline.copy()
    if "overridden_fields" not in result.columns:
        result["overridden_fields"] = [dict() for _ in range(len(result))]
    if "stale_fields" not in result.columns:
        result["stale_fields"] = [dict() for _ in range(len(result))]

    for idx, row in result.iterrows():
        ym_key = f"{int(row['year']):04d}-{int(row['month']):02d}"
        fields = tn_overrides.get(ym_key)
        if not fields:
            continue

        overridden_here: dict = {}
        for field in EDITABLE_MONTH_FIELDS:
            if field in fields:
                overridden_here[field] = result.at[idx, field]  # исходное значение (RSV/прогноз/пусто)
                result.at[idx, field] = fields[field]

        if not overridden_here:
            continue

        result.at[idx, "overridden_fields"] = overridden_here
        result.at[idx, "is_manual_override"] = True

        # Устаревшая правка: сейчас за месяц есть настоящий RSV, а правка вводилась
        # без него (или маркера нет — старая запись) и расходится с тем, что теперь
        # в RSV. Правка по-прежнему применяется (приоритет кадровика), но GUI
        # подсвечивает её, чтобы её не забыли пересмотреть.
        if _flag(row["has_data"]) and fields.get(HAD_DATA_KEY) is not True:
            stale = {
                fld: (orig, fields[fld])
                for fld, orig in overridden_here.items()
                if not _values_equal(orig, fields[fld])
            }
            if stale:
                result.at[idx, "stale_fields"] = stale
        # "(оценка)"/зелёная заливка относятся к СТРОКЕ в целом и означают
        # "это не настоящие данные". Если кадровик поправил только ОДНО поле
        # (например, только kol_rd), остальные поля этого же месяца всё ещё
        # оценочные — снимать пометку со всей строки из-за правки одного
        # поля нельзя, иначе выглядит так, будто и остальные поля тоже стали
        # настоящими данными. Снимаем is_projected только тогда, когда
        # правкой ПОЛНОСТЬЮ перекрыты все редактируемые поля месяца — то
        # есть от исходной "оценки" уже ничего не осталось.
        if set(EDITABLE_MONTH_FIELDS) <= set(overridden_here):
            result.at[idx, "is_projected"] = False
        result.at[idx, "monthly_total"] = (
            _nz(result.at[idx, "vred_dni"]) + _nz(result.at[idx, "otp_dni"])
            + _nz(result.at[idx, "otp_bud_m"]) + _nz(result.at[idx, "po_sredn"])
        )

    # Правка класса на одном месяце должна протянуться вперёд по хронологии
    # так же, как обычный класс из RSV — пересчитываем перенос заново.
    result = _apply_bal_vredn_ffill(result)

    return result


def build_work_year_blocks(
    timeline: pd.DataFrame,
    norm_days: dict[float, float] | None = None,
    month_overrides: dict | None = None,
    termination_date=None,
) -> list[dict]:
    """Группирует таймлайн сотрудника по рабочим годам для отображения.

    Рабочие года, в которых НИ ОДНА строка не имеет has_data=True (обычно —
    "дозаполненный нулями" хвост таймлайна за месяцы, для которых расчёт ещё
    не сделан), в результат вообще не попадают — ни целиком, ни частично.

    Это единственное место, где живёт эта логика — и insalubrity_history.py
    (GUI), и тесты должны использовать именно эту функцию, а не дублировать
    группировку/фильтрацию у себя.

    Возвращает список словарей (по одному на видимый рабочий год, в
    хронологическом порядке), каждый с ключами:
        work_year_key, label, month_rows (список pd.Series — строк таймлайна),
        month_count, has_data, kol_rd_total, vred_total, monthly_total,
        leave (результат summarize_leave_entitlement для этого года, либо None)
    """
    if timeline.empty:
        return []

    timeline = project_current_work_year(timeline, termination_date=termination_date)
    if month_overrides:
        timeline = apply_month_overrides(timeline, month_overrides)

    leave_by_year = {
        entry["work_year_key"]: entry
        for entry in summarize_leave_entitlement(timeline, norm_days=norm_days)
    }

    blocks: dict[tuple, dict] = {}
    order: list[tuple] = []
    for _, r in timeline.iterrows():
        key = r["work_year_key"]
        if key not in blocks:
            blocks[key] = {
                "work_year_key": key,
                "label": r["work_year_label"],
                "month_rows": [],
                "kol_rd_total": 0.0,
                "vred_total": 0.0,
                "monthly_total": 0.0,
                "has_data": False,
                "missing_file_months": [],
            }
            order.append(key)
        b = blocks[key]
        b["month_rows"].append(r)
        b["kol_rd_total"] += _nz(r["kol_rd"])
        b["vred_total"] += _nz(r["vred_dni"])
        b["monthly_total"] += _nz(r["monthly_total"])
        b["has_data"] = b["has_data"] or bool(r["has_data"])
        if _flag(r.get("no_file", False)):
            b["missing_file_months"].append((int(r["year"]), int(r["month"])))

    result = []
    for key in order:
        b = blocks[key]
        if not b["has_data"]:
            continue
        b["month_count"] = len(b["month_rows"])
        b["leave"] = leave_by_year.get(key)
        if b["vred_total"] <= 0:
            # Ни дня реальной вредности за весь рабочий год — значит, класс
            # либо не определён вообще, либо это стухший перенос с давнего
            # года, который мы уже не используем даже для расчёта отпуска
            # (см. проверку vred_total в summarize_leave_entitlement). Не
            # показываем "Балл 3.2" там, где по факту в этом году вредности
            # не было ни дня — это вводило бы в заблуждение.
            for row in b["month_rows"]:
                row["bal_vredn_effective"] = None
                row["bal_vredn_carried"] = False
        result.append(b)
    return result