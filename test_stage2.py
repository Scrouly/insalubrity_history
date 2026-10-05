"""
test_stage2.py — достройка календаря (kol_rd=0 больше не теряет дни вредности).
Запуск: pytest test_stage2.py -v
"""
from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import load_workbook

import insalubrity as ins


def _kalend(rows):
    return pd.DataFrame(rows, columns=["pr_dn", "god", "mes", "kol_rd", "fond_ch"])


def _rows(pr_dn, god, months, kol=22, hours=8.0):
    return [(pr_dn, god, m, kol, kol * hours) for m in months]


def _df(*recs):
    # (tn, pr_dn, god, mes, kol_rd, fond_ch)
    return pd.DataFrame(recs, columns=["tn", "pr_dn", "god", "mes", "kol_rd", "fond_ch"])


def test_nothing_to_fill_is_untouched():
    df = _df((1, 1, 2026, 5, 21.0, 168.0))
    logs = []
    out = ins.fill_missing_calendar(df, _kalend(_rows(1, 2026, [4])), log=logs.append)
    assert out.loc[0, "kol_rd"] == 21 and out.loc[0, "fond_ch"] == 168
    assert logs == []


def test_missing_month_filled_from_previous_months_same_schedule():
    kal = _kalend(_rows(1, 2026, [1, 2, 3, 4], kol=21, hours=8.0) + _rows(2, 2026, [1, 2, 3, 4], kol=20, hours=7.0))
    df = _df((10, 1, 2026, 5, float("nan"), float("nan")))
    logs = []
    out = ins.fill_missing_calendar(df, kal, log=logs.append)
    assert out.loc[0, "kol_rd"] == 21
    assert out.loc[0, "fond_ch"] == pytest.approx(168.0)
    assert round(out.loc[0, "fond_ch"] / out.loc[0, "kol_rd"], 2) == 8.0   # часов в дне сохранены
    assert len(logs) == 1 and "pr_dn=1" in logs[0] and "05.2026" in logs[0] and "10" in logs[0]


def test_zero_in_calendar_itself_is_treated_as_missing():
    kal = _kalend(_rows(1, 2026, [1, 2, 3], kol=22, hours=8.0) + [(1, 2026, 4, 0, 0)])
    df = _df((10, 1, 2026, 4, 0.0, 0.0))
    out = ins.fill_missing_calendar(df, kal, log=lambda m: None)
    assert out.loc[0, "kol_rd"] == 22


def test_uses_following_months_when_no_previous():
    kal = _kalend(_rows(1, 2026, [6, 7], kol=23, hours=8.0))
    df = _df((10, 1, 2026, 5, float("nan"), float("nan")))
    logs = []
    out = ins.fill_missing_calendar(df, kal, log=logs.append)
    assert out.loc[0, "kol_rd"] == 23 and "следующим" in logs[0]


def test_december_label_and_period_math():
    kal = _kalend(_rows(1, 2025, [10, 11], kol=22))
    df = _df((10, 1, 2025, 12, float("nan"), float("nan")))
    logs = []
    out = ins.fill_missing_calendar(df, kal, log=logs.append)
    assert out.loc[0, "kol_rd"] == 22 and "12.2025" in logs[0]


def test_schedule_absent_everywhere_stays_missing_with_error():
    kal = _kalend(_rows(2, 2026, [1, 2, 3]))
    df = _df((10, 1, 2026, 5, float("nan"), float("nan")))
    logs = []
    out = ins.fill_missing_calendar(df, kal, log=logs.append)
    assert pd.isna(out.loc[0, "kol_rd"])
    assert any("ОШИБКА" in m and "pr_dn=1" in m for m in logs)


def test_end_to_end_vredn_days_no_longer_lost(tmp_path, monkeypatch):
    """Через _do_run: месяц без графика в календаре раньше давал vred_dni=0."""
    svod = pd.DataFrame([{
        "tn": 100, "no": 1, "vop": 49, "vf": 176.0, "sux": 176.0 * 0.41, "dolgn": 7,
        "god": 2026, "mes": 5, "pr_dn": 1,
    }])
    lschet = pd.DataFrame([{"tn": 100, "fio": "ИВАНОВ И.И.", "pr_dn": 1}])
    kalend = _kalend(_rows(1, 2026, [1, 2, 3, 4], kol=22, hours=8.0))   # мая нет
    dolgn = pd.DataFrame([{"dshifr": 7, "dname": "Слесарь"}])
    tables = {"lschet": lschet, "kalend": kalend, "dolgn": dolgn, "svod2605": svod}

    monkeypatch.setattr(ins, "find_table", lambda d, name: name)
    monkeypatch.setattr(ins, "load_dbf", lambda path, enc="cp866": tables[path].copy())
    logs = []
    res = ins._do_run(tmp_path, 5, 2026, 0.29, 0.41, 0.58, tmp_path, "cp866", logs.append)

    assert res["sum_rsv"] == 22            # 176 ч / 8 ч в дне
    assert any("нет графика pr_dn=1" in m for m in logs)
    ws = load_workbook(res["rsv_path"]).active
    row = [c.value for c in ws[3]]
    assert row[11] == 22                    # kol_rd (колонка L)
    assert row[13] == 22                    # vred_dni
    assert row[18] == 3.2                   # класс по ставке 0.41
